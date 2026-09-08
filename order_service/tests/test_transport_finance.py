"""Перевозки в финансовом отчёте: расход/приход и лист «Перевозки» (ТЗ 09.2026).

Запуск из папки order_service:  pytest tests/test_transport_finance.py
"""
import io
import os
import sys
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://u:p@localhost/test")
os.environ.setdefault("JWT_SECRET_KEY", "test-secret")
os.environ.setdefault("INTERNAL_API_SECRET", "test-internal-secret")

from app.services import transport_finance  # noqa: E402
from app.services.finance_export import finance_payments_xlsx  # noqa: E402


class _Order:
    def __init__(self, number="п1", ttn="ТТН-2026-Ю000042"):
        self.id = uuid.uuid4()
        self.order_number = number
        self.ttn_number = ttn
        self.delivery_address = "Нефтебаза №1 → База"
        self.status = "delivered"
        self.created_at = datetime(2026, 9, 8, 10, 0, tzinfo=timezone.utc)


class _Detail:
    def __init__(self, **kw):
        self.transport_type = kw.get("transport_type", "to_base")
        self.supplier_payments = kw.get("supplier_payments")
        self.delivery_price = kw.get("delivery_price")
        self.delivery_paid = kw.get("delivery_paid", False)
        self.driver_payment_amount = kw.get("driver_payment_amount")


# ── Расход: оплаты поставщику ────────────────────────────────────────────────

def test_supplier_payments_become_expense_rows():
    detail = _Detail(supplier_payments=[
        {"amount": "20000", "date": "2026-09-01"},
        {"amount": "15000.50", "date": "2026-09-05"},
    ])
    rows = transport_finance.rows_for_order(_Order(), detail)
    assert len(rows) == 2
    assert all(r["direction"] == "expense" for r in rows)
    assert all(r["article"] == transport_finance.ARTICLE_SUPPLIER for r in rows)
    assert rows[1]["amount"] == Decimal("15000.50")
    assert rows[1]["paid_at"] == "2026-09-05"


def test_empty_and_zero_supplier_payments_produce_no_rows():
    detail = _Detail(supplier_payments=[
        {"amount": None, "date": "2026-09-01"}, {"amount": "0", "date": None}, {},
    ])
    assert transport_finance.rows_for_order(_Order(), detail) == []


# ── Приход: стоимость доставки клиенту ───────────────────────────────────────

def test_delivery_price_becomes_income_row():
    detail = _Detail(transport_type="to_client",
                     delivery_price=Decimal("30000"), delivery_paid=True)
    rows = transport_finance.rows_for_order(_Order(), detail)
    assert len(rows) == 1
    assert rows[0]["direction"] == "income"
    assert rows[0]["article"] == transport_finance.ARTICLE_DELIVERY
    assert rows[0]["is_paid"] is True


def test_unpaid_delivery_is_income_but_not_received():
    detail = _Detail(transport_type="to_client",
                     delivery_price=Decimal("30000"), delivery_paid=False)
    rows = transport_finance.rows_for_order(_Order(), detail)
    totals = transport_finance.totals(rows)
    assert totals["income_total"] == Decimal("30000")
    assert totals["income_received"] == Decimal("0")


# ── Оплата водителю — отдельная статья расхода ───────────────────────────────

def test_driver_payment_is_separate_expense_article():
    detail = _Detail(driver_payment_amount=Decimal("12500"))
    rows = transport_finance.rows_for_order(_Order(), detail)
    assert rows[0]["direction"] == "expense"
    assert rows[0]["article"] == transport_finance.ARTICLE_DRIVER


def test_totals_split_supplier_and_driver_expense():
    detail = _Detail(
        supplier_payments=[{"amount": "20000", "date": "2026-09-01"}],
        driver_payment_amount=Decimal("12500"),
        transport_type="to_client",
        delivery_price=Decimal("40000"),
        delivery_paid=True,
    )
    totals = transport_finance.totals(transport_finance.rows_for_order(_Order(), detail))
    assert totals["expense_supplier"] == Decimal("20000")
    assert totals["expense_driver"] == Decimal("12500")
    assert totals["expense_total"] == Decimal("32500")
    assert totals["income_total"] == Decimal("40000")
    assert totals["balance"] == Decimal("7500")


def test_order_without_detail_yields_nothing():
    assert transport_finance.rows_for_order(_Order(), None) == []


def test_build_rows_walks_every_pair():
    pairs = [
        (_Order("п1"), _Detail(driver_payment_amount=Decimal("100"))),
        (_Order("п2"), _Detail(supplier_payments=[{"amount": "200", "date": None}])),
    ]
    rows = transport_finance.build_rows(pairs)
    assert {r["order_number"] for r in rows} == {"п1", "п2"}


# ── Лист «Перевозки» в выгрузке ──────────────────────────────────────────────

def _workbook(transport_rows):
    data = finance_payments_xlsx({
        "period_from": datetime(2026, 9, 1),
        "period_to": datetime(2026, 9, 30),
        "payments": [],
        "transport": transport_rows,
    })
    return load_workbook(io.BytesIO(data))


def test_export_always_has_transport_sheet():
    assert "Перевозки" in _workbook([]).sheetnames


def test_export_marks_every_transport_row():
    rows = transport_finance.rows_for_order(
        _Order(),
        _Detail(supplier_payments=[{"amount": "20000", "date": "2026-09-01"}],
                transport_type="to_client", delivery_price=Decimal("40000")),
    )
    ws = _workbook(rows)["Перевозки"]
    labels = [c.value for c in ws["B"] if c.value == "Перевозка"]
    assert len(labels) == len(rows) == 2


def test_export_puts_expense_and_income_in_separate_columns():
    rows = transport_finance.rows_for_order(
        _Order(),
        _Detail(supplier_payments=[{"amount": "20000", "date": "2026-09-01"}],
                transport_type="to_client", delivery_price=Decimal("40000")),
    )
    ws = _workbook(rows)["Перевозки"]
    values = [(r[7].value, r[8].value) for r in ws.iter_rows()
              if r[1].value == "Перевозка"]
    expense = [v for v in values if v[0] is not None]
    income = [v for v in values if v[1] is not None]
    assert expense == [(20000.0, None)]
    assert income == [(None, 40000.0)]


def test_payments_sheet_is_untouched_by_transport_rows():
    """Регресс: суммы перевозок не должны попасть в лист «Финансы»."""
    rows = transport_finance.rows_for_order(
        _Order(), _Detail(supplier_payments=[{"amount": "999999", "date": None}])
    )
    ws = _workbook(rows)["Финансы"]
    assert not any(
        c.value == 999999.0 for row in ws.iter_rows() for c in row
    )
